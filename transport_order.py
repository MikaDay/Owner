"""Citește detaliile transportului de pe prima pagină a comenzii (format DSV
„Comanda Transport”) și le scrie într-un registru Excel.

Câmpurile care nu se găsesc rămân goale: un PDF cu alt format nu blochează
nimic, doar are mai puține coloane completate.
"""

import io
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from process_pdf import find_order_number

_POSTAL_CITY = re.compile(r"^(?:[A-Z]{1,2}-?\s?)?\d{4,6}\s+(.+)$")  # „310210 Arad”, „D-80331 Munchen”
_GOODS = re.compile(r"^(\d+)\s+(\S+)\s+(.+?)\s+([\d.,]+)\s*kg\s+([\d.,]+)\s*$", re.I)


def _find(pattern, text, flags=re.I):
    m = re.search(pattern, text, flags)
    return m.group(1).strip() if m else ""


def _date(value):
    try:
        return datetime.strptime(value, "%d.%m.%Y").date()
    except (TypeError, ValueError):
        return None


def _number(value):
    """„2.300,00” → 2300.0; „13.60” → 13.6; „20000” → 20000.0"""
    if not value:
        return None
    v = value.replace(" ", "")
    if "," in v:
        v = v.replace(".", "").replace(",", ".")
    try:
        return float(v)
    except ValueError:
        return None


def _stops(lines, kind):
    """Locurile de încărcare/descărcare: [(nume, oraș, dată)], în ordine."""
    stops = []
    header = re.compile(rf"^\d+\.\s*Loc de {kind}", re.I)
    date_re = re.compile(rf"^Data {kind}:\s*(\S+)", re.I)
    for i, line in enumerate(lines):
        if not header.match(line):
            continue
        name = city = date = ""
        for nxt in lines[i + 1:i + 12]:
            m = date_re.match(nxt)
            if m:
                date = m.group(1)
                break
            if not name:
                name = nxt
            elif not city:
                pc = _POSTAL_CITY.match(nxt)
                if pc:
                    city = pc.group(1)
        stops.append((name, city, date))
    return stops


def _place(stops):
    return " → ".join(", ".join(p for p in (name, city.title() if city.isupper() else city) if p)
                      for name, city, _ in stops)


def parse_order(text):
    """Dicționar cu detaliile comenzii din textul primei pagini."""
    lines = [ln.strip() for ln in (text or "").splitlines() if ln.strip()]
    flat = " ".join(lines)

    issuer = ""
    for i, ln in enumerate(lines):
        if re.match(r"^Pagin[aă] 1\s*/", ln, re.I) and i + 1 < len(lines):
            issuer = lines[i + 1]
            break

    loading, unloading = _stops(lines, "incarcare"), _stops(lines, "descarcare")
    goods = weight = ldm = ""
    for i, ln in enumerate(lines):
        if ln.lower().startswith("bucati") and i + 1 < len(lines):
            m = _GOODS.match(lines[i + 1])
            if m:
                goods, weight, ldm = m.group(3), m.group(4), m.group(5)
            break

    order_no = find_order_number(text)
    client_ref = _find(r"Nr\.?\s*Comand[aă]:\s*(\S+)", flat)
    price = re.search(r"Transport:\s*([\d.\s]+,\d{2}|[\d]+(?:\.\d{2})?)\s*([A-Z]{3})", flat)
    return {
        "order_no": order_no,
        "order_date": _date(_find(r"versiune\s*\d+\s*-\s*(\d{2}\.\d{2}\.\d{4})", flat)),
        "issuer": issuer,
        "client_ref": "" if client_ref == order_no else client_ref,
        "dsv_ref": _find(r"DSV Ref\.?:\s*(\S+)", flat),
        "loading": _place(loading),
        "loading_date": _date(loading[0][2]) if loading else None,
        "unloading": _place(unloading),
        "unloading_date": _date(unloading[-1][2]) if unloading else None,
        "goods": goods,
        "weight_kg": _number(weight),
        "ldm": _number(ldm),
        "price": _number(price.group(1)) if price else None,
        "currency": price.group(2) if price else "",
        "truck": _find(r"Numar de inmatriculare camion:\s*(\S+)", flat),
        "contact": _find(r"Nume:\s*(.+?)\s+Telefon:", flat),
    }


COLUMNS = [
    ("Nr. comandă", "order_no", 16, None),
    ("Data comenzii", "order_date", 13, "DD.MM.YYYY"),
    ("Client", "issuer", 22, None),
    ("Loc încărcare", "loading", 34, None),
    ("Data încărcare", "loading_date", 13, "DD.MM.YYYY"),
    ("Loc descărcare", "unloading", 34, None),
    ("Data descărcare", "unloading_date", 14, "DD.MM.YYYY"),
    ("Marfă", "goods", 18, None),
    ("Greutate (kg)", "weight_kg", 13, "#,##0"),
    ("LDM", "ldm", 7, "0.00"),
    ("Preț", "price", 12, "#,##0.00"),
    ("Monedă", "currency", 8, None),
    ("Nr. camion", "truck", 12, None),
    ("Ref. client", "client_ref", 24, None),
    ("Ref. DSV", "dsv_ref", 17, None),
    ("Persoană contact", "contact", 18, None),
    ("Fișier", "file", 28, None),
]


def build_register(rows):
    """Registrul Excel (bytes) pentru o listă de dicționare întoarse de parse_order()."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Transporturi"
    head_fill = PatternFill("solid", fgColor="1D2421")
    for col, (title, _, width, _) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=col, value=title)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 22

    rows = sorted(rows, key=lambda r: (r.get("loading_date") is None, r.get("loading_date") or 0, r.get("order_no", "")))
    for r, row in enumerate(rows, 2):
        for col, (_, key, _, fmt) in enumerate(COLUMNS, 1):
            c = ws.cell(row=r, column=col, value=row.get(key))
            if fmt:
                c.number_format = fmt

    # Total pe monedă, sub tabel (formule, ca să se actualizeze dacă editezi prețurile).
    last = len(rows) + 1
    price_col = get_column_letter([k for _, k, _, _ in COLUMNS].index("price") + 1)
    cur_col = get_column_letter([k for _, k, _, _ in COLUMNS].index("currency") + 1)
    r = last + 2
    for cur in sorted({row.get("currency") for row in rows if row.get("currency")}):
        ws.cell(row=r, column=1, value=f"Total {cur}").font = Font(bold=True)
        c = ws.cell(row=r, column=[k for _, k, _, _ in COLUMNS].index("price") + 1,
                    value=f'=SUMIF({cur_col}2:{cur_col}{last},"{cur}",{price_col}2:{price_col}{last})')
        c.number_format = "#,##0.00"
        c.font = Font(bold=True)
        r += 1

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
